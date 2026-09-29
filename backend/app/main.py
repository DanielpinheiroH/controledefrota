import io
import os
import secrets
import time
from collections import defaultdict, deque
from datetime import date, datetime, timedelta, timezone
from decimal import Decimal
from pathlib import Path
from typing import Annotated
from fastapi import FastAPI, Depends, HTTPException, Query, Request, Response, UploadFile, File, Form
from fastapi.encoders import jsonable_encoder
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, JSONResponse, StreamingResponse, HTMLResponse
from sqlalchemy import select, func, or_, delete, text
from sqlalchemy.exc import IntegrityError
from PIL import Image, UnidentifiedImageError
from pypdf import PdfReader
from .db import get_db
from .models import *
from .schemas import Login, UserInput, TruckCreate, TruckInput, MileageInput, WorkshopInput, MaintenanceInput, PlanInput
from .auth import current_user, admin, password_hash, digest, DUMMY_HASH
from .services import get_or_404, audit, record_mileage, plan_state, add_items
from .uploads import reject_active_pdf

app = FastAPI(title='FrotaGest API', version='1.0.0', docs_url=None, redoc_url=None, openapi_url='/api/openapi.json')
origins = os.getenv('ALLOWED_ORIGINS', 'http://localhost:8088,http://127.0.0.1:8088').split(',')
app.add_middleware(CORSMiddleware, allow_origins=origins, allow_credentials=True, allow_methods=['GET','POST','PUT','DELETE'], allow_headers=['Content-Type','X-Requested-With'])
secure_cookie = os.getenv('COOKIE_SECURE', 'false').lower() == 'true'
attempts = defaultdict(deque)
storage = Path(os.getenv('UPLOAD_DIR', '/app/uploads'))
storage.mkdir(parents=True, exist_ok=True)
Db = Annotated[object, Depends(get_db)]
Reader = Annotated[User, Depends(current_user)]
Writer = Annotated[User, Depends(admin)]

def serialize(obj):
    return {column.name: getattr(obj, column.name) for column in obj.__table__.columns if column.name not in {'password_hash', 'token_hash', 'storage_key'}}

def maintenance_view(db, obj):
    data = serialize(obj)
    data['services'] = [serialize(x) for x in db.scalars(select(MaintenanceService).where(MaintenanceService.maintenance_id == obj.id))]
    data['parts'] = [serialize(x) for x in db.scalars(select(MaintenancePart).where(MaintenancePart.maintenance_id == obj.id))]
    truck = db.get(Truck, obj.truck_id)
    workshop = db.get(Workshop, obj.workshop_id) if obj.workshop_id else None
    data.update(plate=truck.plate, workshop_name=workshop.name if workshop else None)
    return data

def paged(db, query, page, size):
    total = db.scalar(select(func.count()).select_from(query.order_by(None).subquery()))
    return {'items': [serialize(x) for x in db.scalars(query.offset((page-1)*size).limit(size))], 'total': total, 'page': page, 'size': size}

@app.middleware('http')
async def security(request, call_next):
    if request.method in {'POST','PUT','DELETE','PATCH'}:
        origin = request.headers.get('origin')
        if (origin and origin not in origins) or request.headers.get('x-requested-with') != 'FrotaGest':
            return JSONResponse({'detail': 'Origem da requisição inválida'}, status_code=403)
    response = await call_next(request)
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['X-Frame-Options'] = 'DENY'
    response.headers['Referrer-Policy'] = 'same-origin'
    response.headers['Cache-Control'] = 'no-store'
    return response

@app.exception_handler(IntegrityError)
async def conflict(request, exc):
    return JSONResponse({'detail': 'Registro duplicado ou relação inválida. Confira placa, chassi e RENAVAM.'}, status_code=409)

@app.get('/api/health')
def health(db: Db):
    db.execute(text('SELECT 1'))
    return {'status': 'ok'}

@app.get('/api/docs', response_class=HTMLResponse, include_in_schema=False)
def docs():
    return '<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>FrotaGest API</title><link rel="stylesheet" href="/api-docs/swagger-ui.css"></head><body><div id="swagger-ui"></div><script src="/api-docs/swagger-ui-bundle.js"></script><script src="/api-docs/init.js"></script></body></html>'

@app.post('/api/auth/login')
def login(data: Login, request: Request, response: Response, db: Db):
    key = (request.client.host, data.email.lower())
    stamp = time.monotonic()
    queue = attempts[key]
    while queue and queue[0] < stamp - 300:
        queue.popleft()
    if len(queue) >= 10:
        raise HTTPException(429, 'Muitas tentativas. Aguarde cinco minutos.')
    queue.append(stamp)
    user = db.scalar(select(User).where(User.email == data.email.lower()))
    valid = password_hash.verify(data.password, user.password_hash if user else DUMMY_HASH)
    if not user or not valid or not user.active:
        raise HTTPException(401, 'E-mail ou senha inválidos')
    token = secrets.token_urlsafe(48)
    db.execute(delete(Session).where(Session.expires_at < now()))
    db.add(Session(user_id=user.id, token_hash=digest(token), expires_at=now()+timedelta(hours=12)))
    audit(db, user, 'LOGIN', 'user', user.id)
    db.commit()
    response.set_cookie('frotagest_session', token, httponly=True, secure=secure_cookie, samesite='strict', max_age=43200, path='/api')
    return serialize(user)

@app.get('/api/auth/me')
def me(user: Reader):
    return serialize(user)

@app.post('/api/auth/logout')
def logout(request: Request, response: Response, db: Db, user: Reader):
    db.execute(delete(Session).where(Session.token_hash == digest(request.cookies.get('frotagest_session', ''))))
    db.commit()
    response.delete_cookie('frotagest_session', path='/api')
    return {'ok': True}

@app.get('/api/users')
def users(db: Db, user: Writer):
    return [serialize(x) for x in db.scalars(select(User).order_by(User.name))]

@app.post('/api/users', status_code=201)
def create_user(data: UserInput, db: Db, user: Writer):
    obj = User(name=data.name, email=data.email.lower(), role=data.role, password_hash=password_hash.hash(data.password))
    if '@' not in obj.email:
        raise HTTPException(422, 'E-mail inválido')
    db.add(obj)
    db.flush()
    audit(db, user, 'USUARIO_CRIADO', 'user', obj.id)
    db.commit()
    return serialize(obj)

@app.get('/api/trucks')
def trucks(db: Db, user: Reader, q: str = '', active: bool | None = None, page: int = Query(1, ge=1), size: int = Query(30, ge=1, le=100)):
    query = select(Truck).order_by(Truck.active.desc(), Truck.plate)
    if active is not None:
        query = query.where(Truck.active == active)
    if q:
        query = query.where(or_(Truck.plate.ilike(f'%{q}%'), Truck.brand.ilike(f'%{q}%'), Truck.model.ilike(f'%{q}%')))
    return paged(db, query, page, size)

@app.post('/api/trucks', status_code=201)
def create_truck(data: TruckCreate, db: Db, user: Writer):
    obj = Truck(**data.model_dump())
    obj.active = obj.active and obj.status != 'INATIVO'
    if not obj.active:
        obj.status = 'INATIVO'
    db.add(obj)
    db.flush()
    record_mileage(db, obj, obj.mileage, user, source='CADASTRO')
    audit(db, user, 'CAMINHAO_CRIADO', 'truck', obj.id)
    db.commit()
    return serialize(obj)

@app.get('/api/trucks/{ident}')
def truck_detail(ident: int, db: Db, user: Reader):
    return serialize(get_or_404(db, Truck, ident))

@app.put('/api/trucks/{ident}')
def update_truck(ident: int, data: TruckInput, db: Db, user: Writer):
    obj = get_or_404(db, Truck, ident, True)
    before = serialize(obj)
    for key, value in data.model_dump().items():
        setattr(obj, key, value)
    obj.active = obj.active and obj.status != 'INATIVO'
    if not obj.active:
        obj.status = 'INATIVO'
    audit(db, user, 'CAMINHAO_ALTERADO', 'truck', ident, {'before': jsonable_encoder(before), 'after': jsonable_encoder(serialize(obj))})
    db.commit()
    return serialize(obj)

@app.delete('/api/trucks/{ident}')
def deactivate_truck(ident: int, db: Db, user: Writer):
    obj = get_or_404(db, Truck, ident, True)
    obj.active, obj.status = False, 'INATIVO'
    audit(db, user, 'CAMINHAO_INATIVADO', 'truck', ident)
    db.commit()
    return serialize(obj)

@app.post('/api/trucks/{ident}/mileage', status_code=201)
def mileage(ident: int, data: MileageInput, db: Db, user: Writer):
    obj = get_or_404(db, Truck, ident, True)
    record_mileage(db, obj, data.mileage, user, data.notes)
    db.commit()
    return serialize(obj)

@app.get('/api/trucks/{ident}/mileage')
def mileage_history(ident: int, db: Db, user: Reader, page: int = Query(1, ge=1), size: int = Query(30, ge=1, le=100)):
    get_or_404(db, Truck, ident)
    return paged(db, select(MileageHistory).where(MileageHistory.truck_id == ident).order_by(MileageHistory.id.desc()), page, size)

@app.get('/api/workshops')
def workshops(db: Db, user: Reader, q: str = '', page: int = Query(1, ge=1), size: int = Query(30, ge=1, le=100)):
    return paged(db, select(Workshop).where(Workshop.name.ilike(f'%{q}%')).order_by(Workshop.name), page, size)

@app.post('/api/workshops', status_code=201)
def create_workshop(data: WorkshopInput, db: Db, user: Writer):
    obj = Workshop(**data.model_dump())
    db.add(obj)
    db.flush()
    audit(db, user, 'OFICINA_CRIADA', 'workshop', obj.id)
    db.commit()
    return serialize(obj)

@app.put('/api/workshops/{ident}')
def update_workshop(ident: int, data: WorkshopInput, db: Db, user: Writer):
    obj = get_or_404(db, Workshop, ident, True)
    for key, value in data.model_dump().items():
        setattr(obj, key, value)
    audit(db, user, 'OFICINA_ALTERADA', 'workshop', ident)
    db.commit()
    return serialize(obj)

@app.delete('/api/workshops/{ident}')
def deactivate_workshop(ident: int, db: Db, user: Writer):
    obj = get_or_404(db, Workshop, ident, True)
    obj.active = False
    audit(db, user, 'OFICINA_INATIVADA', 'workshop', ident)
    db.commit()
    return serialize(obj)

def maintenance_query(truck_id=None, workshop_id=None, start=None, end=None, status=None):
    query = select(Maintenance).order_by(Maintenance.date.desc(), Maintenance.id.desc())
    for column, value in [(Maintenance.truck_id, truck_id), (Maintenance.workshop_id, workshop_id), (Maintenance.status, status)]:
        if value is not None:
            query = query.where(column == value)
    if start:
        query = query.where(Maintenance.date >= start)
    if end:
        query = query.where(Maintenance.date <= end)
    if start and end and start > end:
        raise HTTPException(422, 'Período inválido')
    return query

@app.get('/api/maintenance')
def maintenance_list(db: Db, user: Reader, truck_id: int | None = None, workshop_id: int | None = None, start: date | None = None, end: date | None = None, status: str | None = None, page: int = Query(1, ge=1), size: int = Query(30, ge=1, le=100)):
    query = maintenance_query(truck_id, workshop_id, start, end, status)
    total = db.scalar(select(func.count()).select_from(query.order_by(None).subquery()))
    return {'items': [maintenance_view(db, x) for x in db.scalars(query.offset((page-1)*size).limit(size))], 'total': total, 'page': page, 'size': size}

def validate_maintenance(db, data):
    truck = get_or_404(db, Truck, data.truck_id, True)
    if data.workshop_id:
        get_or_404(db, Workshop, data.workshop_id)
    if data.plan_id:
        plan = get_or_404(db, MaintenancePlan, data.plan_id)
        if plan.truck_id != data.truck_id or data.type != 'PREVENTIVA':
            raise HTTPException(422, 'Plano incompatível com o caminhão ou tipo')
    if data.status == 'CONCLUIDA' and data.date > date.today():
        raise HTTPException(422, 'Manutenção concluída não pode ter data futura')
    return truck

def resolve_workshop(db, data, user):
    name = (data.workshop_name or '').strip()
    if not name:
        return data.workshop_id
    workshop = db.scalar(select(Workshop).where(func.lower(func.trim(Workshop.name)) == name.lower()).order_by(Workshop.id).limit(1))
    if workshop is None:
        workshop = Workshop(name=name)
        db.add(workshop)
        db.flush()
        audit(db, user, 'OFICINA_CRIADA', 'workshop', workshop.id)
    return workshop.id

@app.post('/api/maintenance', status_code=201)
def create_maintenance(data: MaintenanceInput, db: Db, user: Writer):
    truck = validate_maintenance(db, data)
    fields = data.model_dump(exclude={'services','parts','workshop_name'})
    fields['workshop_id'] = resolve_workshop(db, data, user)
    obj = Maintenance(**fields, created_by=user.id)
    db.add(obj)
    add_items(db, obj, data)
    if obj.status == 'CONCLUIDA' and obj.mileage > truck.mileage:
        record_mileage(db, truck, obj.mileage, user, source='MANUTENCAO')
    audit(db, user, 'MANUTENCAO_CONCLUIDA' if obj.status == 'CONCLUIDA' else 'MANUTENCAO_CRIADA', 'maintenance', obj.id)
    db.commit()
    return maintenance_view(db, obj)

@app.get('/api/maintenance/{ident}')
def maintenance_detail(ident: int, db: Db, user: Reader):
    return maintenance_view(db, get_or_404(db, Maintenance, ident))

@app.put('/api/maintenance/{ident}')
def update_maintenance(ident: int, data: MaintenanceInput, db: Db, user: Writer):
    truck = validate_maintenance(db, data)
    obj = get_or_404(db, Maintenance, ident, True)
    if obj.status in {'CONCLUIDA','CANCELADA'}:
        raise HTTPException(409, 'Registro finalizado: preserve o histórico; cancele e registre uma correção')
    if obj.truck_id != data.truck_id:
        raise HTTPException(422, 'Não é permitido trocar o caminhão do registro')
    before = maintenance_view(db, obj)
    fields = data.model_dump(exclude={'services','parts','workshop_name'})
    fields['workshop_id'] = resolve_workshop(db, data, user)
    for key, value in fields.items():
        setattr(obj, key, value)
    with db.no_autoflush:
        db.execute(delete(MaintenanceService).where(MaintenanceService.maintenance_id == ident))
        db.execute(delete(MaintenancePart).where(MaintenancePart.maintenance_id == ident))
    add_items(db, obj, data)
    if obj.status == 'CONCLUIDA' and obj.mileage > truck.mileage:
        record_mileage(db, truck, obj.mileage, user, source='MANUTENCAO')
    audit(db, user, 'MANUTENCAO_'+('CONCLUIDA' if obj.status == 'CONCLUIDA' else 'ALTERADA'), 'maintenance', ident, {'before': jsonable_encoder(before)})
    db.commit()
    return maintenance_view(db, obj)

@app.delete('/api/maintenance/{ident}')
def cancel_maintenance(ident: int, db: Db, user: Writer):
    obj = get_or_404(db, Maintenance, ident, True)
    obj.status = 'CANCELADA'
    audit(db, user, 'MANUTENCAO_CANCELADA', 'maintenance', ident)
    db.commit()
    return maintenance_view(db, obj)

def plans_data(db, truck_id=None, include_inactive=False):
    query = select(MaintenancePlan).join(Truck)
    if not include_inactive:
        query = query.where(MaintenancePlan.active.is_(True), Truck.active.is_(True))
    if truck_id:
        query = query.where(MaintenancePlan.truck_id == truck_id)
    return [{**serialize(p), **plan_state(db, p, db.get(Truck, p.truck_id))} for p in db.scalars(query.order_by(MaintenancePlan.id))]

@app.get('/api/maintenance-plans')
def plans(db: Db, user: Reader, truck_id: int | None = None):
    return plans_data(db, truck_id, include_inactive=truck_id is not None)

def save_plan(db, data, user, obj=None):
    truck = get_or_404(db, Truck, data.truck_id, True)
    if data.baseline_km > truck.mileage:
        raise HTTPException(422, 'KM da última realização supera o KM atual')
    if obj and obj.truck_id != data.truck_id:
        raise HTTPException(422, 'Não é permitido trocar o caminhão do plano')
    if not obj:
        obj = MaintenancePlan()
        db.add(obj)
    for key, value in data.model_dump(exclude={'last_maintenance_id'}).items():
        setattr(obj, key, value)
    db.flush()
    if data.last_maintenance_id:
        maintenance = get_or_404(db, Maintenance, data.last_maintenance_id, True)
        if maintenance.truck_id != data.truck_id or maintenance.status != 'CONCLUIDA' or maintenance.type != 'PREVENTIVA' or maintenance.plan_id not in (None, obj.id):
            raise HTTPException(422, 'Última manutenção inválida para este plano')
        maintenance.plan_id = obj.id
    audit(db, user, 'PLANO_ALTERADO', 'maintenance_plan', obj.id, jsonable_encoder(data.model_dump()))
    db.commit()
    return {**serialize(obj), **plan_state(db, obj, truck)}

@app.post('/api/maintenance-plans', status_code=201)
def create_plan(data: PlanInput, db: Db, user: Writer):
    return save_plan(db, data, user)

@app.put('/api/maintenance-plans/{ident}')
def update_plan(ident: int, data: PlanInput, db: Db, user: Writer):
    return save_plan(db, data, user, get_or_404(db, MaintenancePlan, ident, True))

@app.get('/api/alerts')
def alerts(db: Db, user: Reader):
    return sorted(plans_data(db), key=lambda p: {'VERMELHO': 0, 'AMARELO': 1, 'VERDE': 2}[p['state']])

def costs_data(db, truck_id=None, workshop_id=None, start=None, end=None):
    rows = list(db.scalars(maintenance_query(truck_id, workshop_id, start, end, 'CONCLUIDA')))
    total = sum((r.total_cost for r in rows), Decimal('0'))
    by_truck, by_workshop, by_month = {}, {}, {}
    for row in rows:
        truck = db.get(Truck, row.truck_id)
        workshop = db.get(Workshop, row.workshop_id) if row.workshop_id else None
        for group, key in [(by_truck, truck.plate), (by_workshop, workshop.name if workshop else 'Sem oficina'), (by_month, row.date.strftime('%Y-%m'))]:
            group[key] = group.get(key, Decimal('0')) + row.total_cost
    return {'total': total, 'parts': sum((r.parts_cost for r in rows), Decimal('0')), 'labor': sum((r.labor_cost for r in rows), Decimal('0')), 'other': sum((r.other_cost for r in rows), Decimal('0')), 'by_truck': by_truck, 'by_workshop': by_workshop, 'by_month': by_month}

@app.get('/api/reports/costs')
def costs(db: Db, user: Reader, truck_id: int | None = None, workshop_id: int | None = None, start: date | None = None, end: date | None = None):
    return costs_data(db, truck_id, workshop_id, start, end)

@app.get('/api/dashboard')
def dashboard(db: Db, user: Reader):
    trucks = list(db.scalars(select(Truck).where(Truck.active.is_(True))))
    plans = plans_data(db)
    today = date.today()
    return {'total': len(trucks), 'available': sum(t.status == 'DISPONIVEL' for t in trucks), 'maintenance': sum(t.status == 'EM_MANUTENCAO' for t in trucks), 'stopped': sum(t.status == 'PARADO' for t in trucks), 'upcoming': sum(p['state'] == 'AMARELO' for p in plans), 'overdue': sum(p['state'] == 'VERMELHO' for p in plans), 'month_cost': costs_data(db, start=today.replace(day=1), end=today)['total'], 'year_cost': costs_data(db, start=today.replace(month=1, day=1), end=today)['total'], 'alerts': sorted([p for p in plans if p['state'] != 'VERDE'], key=lambda p: p['state'] != 'VERMELHO'), 'recent': [maintenance_view(db, m) for m in db.scalars(maintenance_query().limit(5))]}

@app.get('/api/reports/export.xlsx')
def export(db: Db, user: Reader, truck_id: int | None = None, start: date | None = None, end: date | None = None):
    from openpyxl import Workbook
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Manutenções'
    sheet.append(['Data','Placa','Descrição','Status','KM','Peças','Mão de obra','Outros','Total'])
    for row in db.scalars(maintenance_query(truck_id, start=start, end=end).limit(10000)):
        sheet.append([row.date, db.get(Truck, row.truck_id).plate, row.description, row.status, row.mileage, row.parts_cost, row.labor_cost, row.other_cost, row.total_cost])
        for cell in sheet[sheet.max_row]:
            if cell.data_type == 'f' or isinstance(cell.value, str):
                cell.data_type = 's'
        sheet.cell(sheet.max_row, 1).number_format = 'DD/MM/YYYY'
        for col in range(6,10):
            sheet.cell(sheet.max_row, col).number_format = '"R$" #,##0.00'
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    for col in 'ABCDEFGHI':
        sheet.column_dimensions[col].width = 24 if col != 'C' else 45
    buffer = io.BytesIO()
    workbook.save(buffer)
    buffer.seek(0)
    return StreamingResponse(buffer, media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', headers={'Content-Disposition': 'attachment; filename=frotagest.xlsx'})

@app.get('/api/trucks/{ident}/attachments')
def attachments(ident: int, db: Db, user: Reader, maintenance_id: int | None = None):
    get_or_404(db, Truck, ident)
    query = select(Attachment).where(Attachment.truck_id == ident)
    if maintenance_id is not None:
        query = query.where(Attachment.maintenance_id == maintenance_id)
    return [serialize(a) for a in db.scalars(query.order_by(Attachment.id.desc()))]

@app.post('/api/trucks/{ident}/attachments', status_code=201)
async def upload(ident: int, db: Db, user: Writer, file: UploadFile = File(...), kind: str = Form('DOCUMENTO'), maintenance_id: int | None = Form(None)):
    get_or_404(db, Truck, ident)
    if maintenance_id and get_or_404(db, Maintenance, maintenance_id).truck_id != ident:
        raise HTTPException(422, 'Manutenção pertence a outro caminhão')
    if kind not in {'DOCUMENTO','FOTO'}:
        raise HTTPException(422, 'Tipo inválido')
    content = await file.read(10*1024*1024+1)
    if len(content) > 10*1024*1024:
        raise HTTPException(413, 'Limite de 10 MB por arquivo')
    ext = Path(file.filename or '').suffix.lower()
    mime = None
    try:
        if ext == '.pdf' and file.content_type == 'application/pdf' and content.startswith(b'%PDF-') and kind != 'FOTO':
            pdf = PdfReader(io.BytesIO(content))
            if pdf.is_encrypted:
                raise ValueError('PDF ativo ou criptografado')
            reject_active_pdf(pdf)
            if not len(pdf.pages):
                raise ValueError('PDF vazio')
            mime = 'application/pdf'
        elif ext in {'.jpg','.jpeg','.png','.webp'} and file.content_type in {'image/jpeg','image/png','image/webp'}:
            picture = Image.open(io.BytesIO(content))
            if picture.format not in {'JPEG','PNG','WEBP'} or picture.width * picture.height > 25000000:
                raise ValueError('Imagem inválida')
            picture.load()
            output = io.BytesIO()
            picture.convert('RGB').save(output, format='JPEG', quality=90)
            content, ext, mime = output.getvalue(), '.jpg', 'image/jpeg'
        else:
            raise ValueError('Formato não permitido')
    except Exception:
        raise HTTPException(422, 'Arquivo inválido. Envie PDF sem conteúdo ativo ou imagem JPG, PNG ou WebP.')
    key = secrets.token_hex(24)+ext
    target = storage/key
    target.write_bytes(content)
    try:
        obj = Attachment(truck_id=ident, maintenance_id=maintenance_id, filename=Path(file.filename or 'documento').name[:200], storage_key=key, mime=mime, kind=kind, size=len(content))
        db.add(obj)
        db.flush()
        audit(db, user, 'DOCUMENTO_ANEXADO', 'attachment', obj.id)
        db.commit()
    except Exception:
        target.unlink(missing_ok=True)
        raise
    return serialize(obj)

@app.get('/api/attachments/{ident}/download')
def download(ident: int, db: Db, user: Reader):
    obj = get_or_404(db, Attachment, ident)
    return FileResponse(storage/obj.storage_key, media_type=obj.mime, filename=obj.filename, content_disposition_type='inline' if obj.mime.startswith('image/') else 'attachment')

@app.get('/api/audit')
def audit_list(db: Db, user: Writer, page: int = Query(1, ge=1), size: int = Query(30, ge=1, le=100)):
    return paged(db, select(AuditLog).order_by(AuditLog.id.desc()), page, size)
