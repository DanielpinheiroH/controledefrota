"""Exercise company boundaries through the public API, in both directions."""
from datetime import date
import io
import pytest
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from app.db import get_db
from app.main import app
from app.models import MileageHistory, Tenant, User
from app.tenancy import scoped
from test_flows import truck, maintenance, plan


def login(client, tenant):
    response = client.post('/api/auth/login', json={
        'tenant_id': tenant, 'email': 'admin@example.test',
        'password': 'TenantOnePassword!' if tenant == 1 else 'TestingPassword123!'})
    assert response.status_code == 200, response.text
    assert response.json()['tenant_id'] == tenant


@pytest.mark.parametrize('first,second', [(3, 1), (1, 3)])
def test_company_boundaries(client, first, second):
    login(client, first)
    t = truck(client)
    w = client.post('/api/workshops', json={'name': 'Oficina isolada'}).json()
    m = maintenance(client, t['id'], workshop_id=w['id'], parts=[{'name':'Filtro','quantity':1,'unit_price':'40.00'}])
    p = plan(client, t['id'], last_maintenance_id=m['id'])
    photo = io.BytesIO()
    Image.new('RGB', (2, 2)).save(photo, format='PNG')
    a = client.post(f"/api/trucks/{t['id']}/attachments", files={'file':('foto.png', photo.getvalue(), 'image/png')}).json()
    assert client.get(f"/api/attachments/{a['id']}/download").status_code == 200
    login(client, second)
    assert client.get('/api/trucks', headers={'X-FrotaGest-Tenant':str(first)}).status_code == 401
    assert client.post('/api/workshops', json={'name':'Aba antiga'}, headers={'X-FrotaGest-Tenant':str(first)}).status_code == 401
    # A request header/query cannot switch the authenticated company.
    for endpoint in ('trucks', 'workshops', 'maintenance'):
        response = client.get(f'/api/{endpoint}?tenant_id={first}', headers={'X-Tenant-ID': str(first)})
        assert response.json()['total'] == 0
        assert response.json()['items'] == []
    for endpoint in ('maintenance-plans', 'alerts'):
        assert client.get(f'/api/{endpoint}').json() == []
    dashboard = client.get('/api/dashboard').json()
    assert dashboard['total'] == dashboard['month_cost'] == dashboard['year_cost'] == 0
    assert dashboard['recent'] == dashboard['alerts'] == []
    assert client.get('/api/reports/costs').json()['total'] == 0
    workbook = load_workbook(io.BytesIO(client.get('/api/reports/export.xlsx').content))
    assert workbook.active.max_row == 1
    assert all(u['tenant_id'] == second for u in client.get('/api/users').json())
    assert all(a['tenant_id'] == second for a in client.get('/api/audit').json()['items'])
    for path in (f"trucks/{t['id']}", f"trucks/{t['id']}/mileage", f"trucks/{t['id']}/attachments",
                 f"maintenance/{m['id']}", f"attachments/{a['id']}/download"):
        assert client.get('/api/' + path).status_code == 404
    for path in (f"trucks/{t['id']}", f"workshops/{w['id']}", f"maintenance/{m['id']}"):
        assert client.delete('/api/' + path).status_code == 404
    assert client.post(f"/api/trucks/{t['id']}/mileage", json={'mileage':400000}).status_code == 404
    assert client.post(f"/api/trucks/{t['id']}/attachments", files={'file':('f.png', photo.getvalue(),'image/png')}).status_code == 404
    own = truck(client)  # Same plate is valid in another company.
    data = {'truck_id':own['id'], 'type':'PREVENTIVA','category':'Motor','description':'Isolamento',
            'date':str(date.today()),'mileage':325000,'status':'ABERTA'}
    for fields in ({'truck_id':t['id']}, {'workshop_id':w['id']}, {'plan_id':p['id']}):
        assert client.post('/api/maintenance', json={**data, **fields}).status_code == 404
    assert client.put(f"/api/maintenance/{m['id']}", json=data).status_code == 404
    assert client.put(f"/api/trucks/{t['id']}", json={'plate':'XYZ9A99','brand':'A','model':'B'}).status_code == 404
    assert client.put(f"/api/workshops/{w['id']}", json={'name':'Alterada'}).status_code == 404
    pdata = {'truck_id':own['id'],'name':'Plano','interval_km':1000,'baseline_km':100,'baseline_date':str(date.today())}
    assert client.put(f"/api/maintenance-plans/{p['id']}", json=pdata).status_code == 404
    assert client.post('/api/maintenance-plans', json={**pdata,'last_maintenance_id':m['id']}).status_code == 404
    assert client.post(f"/api/trucks/{own['id']}/attachments", data={'maintenance_id':m['id']}, files={'file':('f.png',photo.getvalue(),'image/png')}).status_code == 404
    assert client.post('/api/users',json={'tenant_id':first,'name':'Attack','email':'a@b.test','password':'TestingPassword123!'}).status_code == 422
    own_m = maintenance(client, own['id'], workshop_name='Oficina isolada')
    assert own_m['workshop_id'] != w['id']
    assert own_m['tenant_id'] == second
    login(client, first)
    assert client.get('/api/trucks').json()['total'] == 1
    assert client.get(f"/api/trucks/{t['id']}").json()['active'] is True
    assert client.get(f"/api/maintenance/{m['id']}").json()['status'] == 'CONCLUIDA'
    assert client.get(f"/api/maintenance/{own_m['id']}").status_code == 404


def test_login_and_database_boundaries(client):
    assert client.post('/api/auth/login',json={'tenant_id':1,'email':'admin@example.test','password':'TestingPassword123!'}).status_code == 401
    assert client.post('/api/auth/login',json={'tenant_id':999,'email':'admin@example.test','password':'TestingPassword123!'}).status_code == 401
    login(client, 3)
    t = truck(client)
    dependency = app.dependency_overrides[get_db]()
    db = next(dependency)
    try:
        with pytest.raises(RuntimeError):
            scoped(db, User)
        other = db.scalar(select(User).where(User.tenant_id == 1))
        with pytest.raises(IntegrityError), db.begin_nested():
            db.add(MileageHistory(tenant_id=1, truck_id=t['id'], user_id=other.id, mileage=1))
            db.flush()
        db.get(Tenant, 3).active = False
        db.commit()
        assert client.get('/api/auth/me').status_code == 401
    finally:
        dependency.close()
