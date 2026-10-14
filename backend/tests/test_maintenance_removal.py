from datetime import date, timedelta
import io
from PIL import Image
from openpyxl import load_workbook
from test_flows import truck, maintenance, plan


def remove(client, ident, code=200):
    result = client.post(f'/api/maintenance/{ident}/remove', json={})
    assert result.status_code == code, result.text
    return result.json()


def test_remove_last_restores_km_costs_plan_and_hides_attachments(admin):
    t = truck(admin, mileage=200000)
    p = plan(admin, t['id'], baseline_km=200000)
    first = maintenance(admin, t['id'], 210000, plan_id=p['id'])
    last = maintenance(admin, t['id'], 220000, plan_id=p['id'])
    assert not admin.get(f"/api/maintenance/{first['id']}").json()['can_delete']
    assert admin.get(f"/api/maintenance/{last['id']}").json()['can_delete']
    photo = io.BytesIO()
    Image.new('RGB', (2,2)).save(photo, format='PNG')
    attachment = admin.post(f"/api/trucks/{t['id']}/attachments", data={'maintenance_id':last['id']}, files={'file':('test.png',photo.getvalue(),'image/png')}).json()
    remove(admin, first['id'], 409)
    assert remove(admin, last['id'])['mileage'] == 210000
    assert admin.get(f"/api/trucks/{t['id']}").json()['mileage'] == 210000
    assert admin.get(f"/api/maintenance/{last['id']}").status_code == 404
    assert admin.get('/api/maintenance').json()['total'] == 1
    assert admin.get('/api/reports/costs').json()['total'] == 1200
    assert admin.get('/api/alerts').json()[0]['next_km'] == 225000
    assert len(admin.get('/api/dashboard').json()['recent']) == 1
    assert admin.get(f"/api/attachments/{attachment['id']}/download").status_code == 404
    assert admin.get(f"/api/trucks/{t['id']}/attachments").json() == []
    assert load_workbook(io.BytesIO(admin.get('/api/reports/export.xlsx').content)).active.max_row == 2
    assert admin.get(f"/api/trucks/{t['id']}/mileage").json()['items'][0]['reverted'] is True
    assert any(a['action']=='MANUTENCAO_EXCLUIDA' for a in admin.get('/api/audit').json()['items'])
    remove(admin, last['id'], 404)
    assert remove(admin, first['id'])['mileage'] == 200000
    assert admin.get('/api/reports/costs').json()['total'] == 0


def test_later_manual_km_is_preserved_through_sequential_removals(admin):
    t = truck(admin, mileage=200000)
    first = maintenance(admin, t['id'], 210000)
    admin.post(f"/api/trucks/{t['id']}/mileage", json={'mileage':215000})
    last = maintenance(admin, t['id'], 220000)
    assert remove(admin, last['id'])['mileage'] == 215000
    assert remove(admin, first['id'])['mileage'] == 215000


def test_same_km_or_open_maintenance_does_not_undo_another_event(admin):
    t = truck(admin, mileage=200000)
    first = maintenance(admin, t['id'], 210000)
    equal = maintenance(admin, t['id'], 210000)
    pending = maintenance(admin, t['id'], 250000, status='ABERTA')
    assert remove(admin, pending['id'])['mileage'] == 210000
    assert remove(admin, equal['id'])['mileage'] == 210000
    assert remove(admin, first['id'])['mileage'] == 200000


def test_manual_km_after_the_last_maintenance_is_not_reduced(admin):
    t = truck(admin, mileage=200000)
    m = maintenance(admin, t['id'], 210000)
    admin.post(f"/api/trucks/{t['id']}/mileage", json={'mileage':215000})
    result = remove(admin, m['id'])
    assert result['mileage'] == 215000
    assert result['mileage_restored'] is False


def test_completion_on_edit_can_be_removed(admin):
    t = truck(admin, mileage=200000)
    m = maintenance(admin, t['id'], 210000, status='ABERTA')
    payload = {key:m[key] for key in ('truck_id','type','category','description','date','mileage')}
    assert admin.put(f"/api/maintenance/{m['id']}", json={**payload,'status':'CONCLUIDA'}).status_code == 200
    assert remove(admin, m['id'])['mileage'] == 200000


def test_removal_requires_admin_and_same_company(admin):
    t = truck(admin, mileage=200000)
    m = maintenance(admin, t['id'], 210000)
    admin.post('/api/auth/login', json={'email':'reader@example.test','password':'TestingPassword123!'})
    remove(admin, m['id'], 403)
    admin.post('/api/auth/login', json={'tenant_id':1,'email':'admin@example.test','password':'TenantOnePassword!'})
    remove(admin, m['id'], 404)


def test_dates_are_inclusive_and_combine_with_truck_filter(admin):
    a, b = truck(admin), truck(admin, plate='XYZ9A99')
    start, end = date.today()-timedelta(days=10), date.today()-timedelta(days=2)
    before = maintenance(admin,a['id'],date=str(start-timedelta(days=1)))
    first = maintenance(admin,a['id'],date=str(start))
    last = maintenance(admin,a['id'],date=str(end))
    maintenance(admin,a['id'],date=str(end+timedelta(days=1)))
    maintenance(admin,b['id'],date=str(start))
    result = admin.get(f"/api/maintenance?truck_id={a['id']}&start={start}&end={end}").json()
    assert result['total'] == 2
    assert [m['id'] for m in result['items']] == [last['id'],first['id']]
    remove(admin,before['id'],409)
    assert admin.get(f'/api/maintenance?start={end}&end={start}').status_code == 422
